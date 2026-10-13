"""Non-GUI tests; can run without PySide6 or internet."""
from __future__ import annotations

import csv
import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import Workbook

from glassbalance.database import Ledger
from glassbalance.email_backup import decrypt_snapshot, encrypt_snapshot
from glassbalance.importer import load_import
from glassbalance.money import as_mills, display_money


class MoneyTests(unittest.TestCase):
    def test_exact_fractional_cent(self):
        self.assertEqual(as_mills("-0.735"),-735)
        self.assertEqual(as_mills("-107.705"),-107705)
        self.assertEqual(as_mills("-0.735")+as_mills("-107.705"),-108440)
        self.assertEqual(display_money(-735),"−$0.735")
        self.assertEqual(display_money(3522000),"$3,522.00")
        with self.assertRaises(ValueError):
            as_mills("NaN")
        with self.assertRaises(ValueError):
            as_mills("1.2345")


class LedgerTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.db=Ledger(Path(self.tmp.name)/"ledger.sqlite")

    def tearDown(self):
        self.db.close()
        self.tmp.cleanup()

    def test_crud_balance_and_date_order(self):
        self.db.set_opening(390420,"2022-08-30","estimated")
        later=self.db.add("2022-09-02","Rent",-672000)
        earlier=self.db.add("2022-08-31","August Salary",1204550)
        self.assertEqual(self.db.balance(),922970)
        entries=self.db.entries()
        self.assertEqual([e.description for e in entries],["August Salary","Rent"])
        self.assertEqual([e.running_mills for e in entries],[1594970,922970])
        self.db.edit(later,"2022-09-03","Rent",-600000)
        self.assertEqual(self.db.balance(),994970)
        self.db.delete(earlier)
        self.assertEqual(self.db.balance(),-209580)
        self.assertEqual(self.db.count(),1)

    def test_roundtrip_export_import_and_duplicates(self):
        self.db.set_opening(100000,"2022-08-30","estimated")
        self.db.add("2022-08-31","August Salary",1100550,"estimated")
        self.db.add("2022-09-02","Rent",-700735,"confirmed")
        csv_file=Path(self.tmp.name)/"ledger.csv"
        xlsx_file=Path(self.tmp.name)/"ledger.xlsx"
        self.db.export_csv(csv_file)
        self.db.export_xlsx(xlsx_file)
        for file in (csv_file,xlsx_file):
            data=load_import(file)
            self.assertEqual(data.opening_mills,100000)
            self.assertEqual(len(data.transactions),2)
            self.assertEqual(sum(t.mills for t in data.transactions),399815)
            with _temporary_ledger() as another:
                self.assertEqual(another.import_data(data),(2,0))
                self.assertEqual(another.balance(),499815)
                self.assertEqual(another.import_data(data),(0,2))

    def test_sqlite_backup_and_restore(self):
        self.db.set_opening(300000,"2023-01-01")
        self.db.add("2023-01-02","Meal",-10500)
        saved=self.db.backup_bytes()
        self.assertTrue(saved.startswith(b"SQLite format 3\0"))
        self.db.add("2023-01-03","Taxi",-5000)
        self.assertEqual(self.db.count(),2)
        self.db.restore_bytes(saved)
        self.assertEqual(self.db.count(),1)
        self.assertEqual(self.db.balance(),289500)

    def test_monthly_summary(self):
        self.db.set_opening(100000,"2026-08-30")
        self.db.add("2026-08-31","Salary",210000)
        self.db.add("2026-09-02","Rent",-80000)
        months=self.db.monthly_totals()
        self.assertEqual([m["month"] for m in months],["2026-08","2026-09"])
        self.assertEqual(months[0]["closing"],310000)
        self.assertEqual(months[1]["closing"],230000)


from contextlib import contextmanager
@contextmanager
def _temporary_ledger():
    db=Ledger(":memory:")
    try:
        yield db
    finally:
        db.close()


class SourceImportTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.path=Path(self.tmp.name)/"original.xlsx"
        wb=Workbook()
        ws=wb.active
        ws.title="Sheet1"
        ws["A1"]="Name"
        ws["B1"]="Bill or Income"
        ws["C2"]=390.42
        ws["E2"]="August 2022 - October 2022"
        for row,(name,amount) in enumerate((
            ("August Salary",1204.55),("lunch",-5),
            ("house rent",-672),("green pepper",-0.735),
            ("Sept Salary",1600),("rent",-400),
            ("kub",-107.705)
        ),start=3):
            ws.cell(row,1,name)
            ws.cell(row,2,amount)
        wb.save(self.path)

    def tearDown(self):
        self.tmp.cleanup()

    def test_original_dates_and_exact_amounts(self):
        data=load_import(self.path,date(2022,10,12))
        self.assertEqual(data.opening_mills,390420)
        self.assertEqual(data.opening_date,"2022-08-30")
        self.assertEqual(len(data.transactions),7)
        by_desc={tx.description:tx for tx in data.transactions}
        self.assertEqual(by_desc["August Salary"].date,"2022-08-31")
        self.assertEqual(by_desc["house rent"].date,"2022-09-02")
        self.assertEqual(by_desc["Sept Salary"].date,"2022-09-30")
        self.assertEqual(by_desc["rent"].date,"2022-10-02")
        self.assertEqual(by_desc["green pepper"].mills,-735)
        self.assertTrue(all(tx.date_status=="estimated" for tx in data.transactions))
        self.assertTrue(all(tx.source_cell for tx in data.transactions))
        with _temporary_ledger() as ledger:
            ledger.import_data(data)
            expected=390420+1204550-5000-672000-735+1600000-400000-107705
            self.assertEqual(ledger.balance(),expected)

    def test_original_repeat_from_csv_provenance(self):
        data=load_import(self.path,date(2022,10,12))
        with _temporary_ledger() as ledger:
            self.assertEqual(ledger.import_data(data),(7,0))
            export=Path(self.tmp.name)/"export.csv"
            ledger.export_csv(export)
            self.assertEqual(ledger.import_data(load_import(export)),(0,7))


class EncryptionTests(unittest.TestCase):
    def test_encrypted_backup_roundtrip_and_tamper(self):
        with _temporary_ledger() as db:
            db.add("2026-10-12","Lunch",-12340)
            snapshot=db.backup_bytes()
        secret="correct horse battery staple"
        encrypted=encrypt_snapshot(snapshot,secret)
        self.assertNotIn(b"Lunch",encrypted)
        self.assertEqual(decrypt_snapshot(encrypted,secret),snapshot)
        with self.assertRaises(ValueError):
            decrypt_snapshot(encrypted,"not the right password")
        altered=bytearray(encrypted)
        altered[-5]^=1
        with self.assertRaises(ValueError):
            decrypt_snapshot(bytes(altered),secret)
        with self.assertRaises(ValueError):
            encrypt_snapshot(snapshot,"short")


if __name__=="__main__":
    unittest.main()
