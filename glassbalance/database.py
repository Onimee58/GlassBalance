"""Small SQLite ledger. No transaction balances are stored or trusted."""
from __future__ import annotations

import csv
import sqlite3
import tempfile
from dataclasses import dataclass
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Iterable
from uuid import uuid4

from .money import mills_as_decimal
from . import comments, categories, recurring, unbalanced


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def clean_date(value: str) -> str:
    return date.fromisoformat(str(value)).isoformat()


@dataclass(frozen=True)
class Entry:
    id: int
    date: str
    description: str
    mills: int
    date_status: str
    source_uid: str
    source_period: str
    source_cell: str
    source_order: int | None
    running_mills: int


class Ledger:
    def __init__(self, path: str | Path = ":memory:"):
        self.path = Path(path) if str(path) != ":memory:" else None
        if self.path:
            self.path.parent.mkdir(parents=True, exist_ok=True)
        self.conn = sqlite3.connect(str(path))
        self.conn.row_factory = sqlite3.Row
        self.conn.execute("PRAGMA foreign_keys=ON")
        self.conn.execute("PRAGMA busy_timeout=5000")
        if self.path:
            self.conn.execute("PRAGMA journal_mode=WAL")
        self.conn.executescript("""
            CREATE TABLE IF NOT EXISTS state (
                id INTEGER PRIMARY KEY CHECK (id=1),
                opening_mills INTEGER NOT NULL DEFAULT 0,
                opening_date TEXT NOT NULL DEFAULT '1970-01-01',
                opening_status TEXT NOT NULL DEFAULT 'confirmed'
            );
            INSERT OR IGNORE INTO state(id, opening_mills, opening_date, opening_status)
                VALUES (1, 0, '1970-01-01', 'confirmed');
            CREATE TABLE IF NOT EXISTS transactions (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                tx_date TEXT NOT NULL,
                description TEXT NOT NULL,
                amount_mills INTEGER NOT NULL,
                date_status TEXT NOT NULL CHECK(date_status IN ('estimated','confirmed')),
                source_uid TEXT UNIQUE NOT NULL,
                source_period TEXT NOT NULL DEFAULT '',
                source_cell TEXT NOT NULL DEFAULT '',
                source_order INTEGER,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL
            );
            CREATE INDEX IF NOT EXISTS idx_tx_date ON transactions(tx_date);
        """)

        self._install_extensions()

    def _install_extensions(self) -> None:
        """Create optional feature tables. Idempotent; never touches existing rows."""
        comments.install(self.conn)
        unbalanced.install(self)
        categories.install(self.conn)
        recurring.install(self.conn)

    def category(self, tx_id: int) -> str:
        return categories.get(self.conn, tx_id)

    def set_category(self, tx_id: int, name: str | None) -> str:
        return categories.set_category(self.conn, tx_id, name)

    def loan_linked(self, tx_id: int) -> bool:
        return self.conn.execute("""SELECT 1 FROM loan_records WHERE origin_tx_id=?
            UNION SELECT 1 FROM loan_repayments WHERE repayment_tx_id=?""",(tx_id,tx_id)).fetchone() is not None

    def comment(self, tx_id: int) -> str:
        return comments.get(self.conn, tx_id)

    def set_comment(self, tx_id: int, note: str) -> None:
        comments.set_note(self.conn, tx_id, note)

    def close(self) -> None:
        self.conn.close()

    def opening(self) -> tuple[int, str, str]:
        row = self.conn.execute("SELECT opening_mills,opening_date,opening_status FROM state WHERE id=1").fetchone()
        return int(row["opening_mills"]), row["opening_date"], row["opening_status"]

    def set_opening(self, mills: int, dt: str, status: str = "confirmed") -> None:
        if status not in ("estimated", "confirmed"):
            raise ValueError("Unknown date status")
        with self.conn:
            self.conn.execute("UPDATE state SET opening_mills=?,opening_date=?,opening_status=? WHERE id=1",
                              (int(mills), clean_date(dt), status))

    def count(self) -> int:
        return int(self.conn.execute("SELECT COUNT(*) FROM transactions").fetchone()[0])

    def _check_entry(self, dt: str, description: str, mills: int, status: str) -> tuple[str,str,int,str]:
        dt = clean_date(dt)
        description = description.strip()
        if not description or len(description)>300:
            raise ValueError("Description must contain 1 to 300 characters.")
        if not isinstance(mills, int) or isinstance(mills, bool):
            raise ValueError("Amount must be represented as integer mills.")
        if abs(mills)>100_000_000_000:
            raise ValueError("Amount is outside the supported range.")
        if status not in ("estimated", "confirmed"):
            raise ValueError("Unknown date status")
        return dt,description,mills,status

    def add(self, dt: str, description: str, mills: int, status: str = "confirmed",
            source_uid: str | None = None, source_period: str = "",
            source_cell: str = "", source_order: int | None = None) -> int:
        dt,description,mills,status=self._check_entry(dt,description,mills,status)
        now=utc_now()
        source_uid=source_uid or f"glassbalance:manual:{uuid4()}"
        with self.conn:
            cur=self.conn.execute("""INSERT INTO transactions(
                tx_date,description,amount_mills,date_status,source_uid,source_period,
                source_cell,source_order,created_at,updated_at
            ) VALUES(?,?,?,?,?,?,?,?,?,?)""",
            (dt,description,mills,status,source_uid,source_period,source_cell,source_order,now,now))
        return int(cur.lastrowid)

    def edit(self, entry_id: int, dt: str, description: str, mills: int,
             status: str="confirmed") -> None:
        dt,description,mills,status=self._check_entry(dt,description,mills,status)
        # Preserve confirmed loan relationships: unlink before changing linked amounts.
        original=self.conn.execute("SELECT amount_mills FROM transactions WHERE id=?",(entry_id,)).fetchone()
        if original and int(original[0]) != mills:
            loan_tables={row[0] for row in self.conn.execute("SELECT name FROM sqlite_master WHERE type='table'")}
            if 'loan_records' in loan_tables:
                linked=self.conn.execute("""SELECT 1 FROM loan_records WHERE origin_tx_id=?
                    UNION SELECT 1 FROM loan_repayments WHERE repayment_tx_id=?""",(entry_id,entry_id)).fetchone()
                if linked:
                    raise ValueError("This transaction is linked to a loan. Unlink or remove its loan tracking before changing the amount.")
        with self.conn:
            cur=self.conn.execute("""UPDATE transactions SET tx_date=?,description=?,
                amount_mills=?,date_status=?,updated_at=? WHERE id=?""",
                (dt,description,mills,status,utc_now(),entry_id))
            if cur.rowcount!=1:
                raise ValueError("The selected transaction no longer exists.")

    def delete(self, entry_id: int) -> dict:
        """Delete one transaction and return a snapshot that restore_deleted() accepts.

        Loan-linked transactions are refused: deleting them would cascade away loan
        history that an undo could not faithfully rebuild.
        """
        if self.loan_linked(entry_id):
            raise ValueError("This transaction is tracked in Unbalanced. Unlink the repayment or remove "
                             "the loan tracking before deleting it.")
        row=self.conn.execute("SELECT * FROM transactions WHERE id=?",(entry_id,)).fetchone()
        if row is None:
            raise ValueError("The selected transaction no longer exists.")
        snapshot={
            "transaction": dict(row),
            "comment": comments.get(self.conn, entry_id),
            "category": categories.get(self.conn, entry_id),
            "recurring_posts": [dict(r) for r in self.conn.execute(
                "SELECT template_id,month FROM recurring_posts WHERE transaction_id=?",(entry_id,))],
            "dismissed": [int(r[0]) for r in self.conn.execute(
                "SELECT loan_id FROM dismissed_suggestions WHERE repayment_tx_id=?",(entry_id,))],
        }
        with self.conn:
            cur=self.conn.execute("DELETE FROM transactions WHERE id=?",(entry_id,))
            if cur.rowcount!=1:
                raise ValueError("The selected transaction no longer exists.")
        return snapshot

    def restore_deleted(self, snapshot: dict) -> int:
        """Re-insert a deleted transaction with its original ID, comment and category."""
        tx=snapshot["transaction"]
        if self.conn.execute("SELECT 1 FROM transactions WHERE id=? OR source_uid=?",
                             (tx["id"],tx["source_uid"])).fetchone():
            raise ValueError("This transaction already exists; nothing was restored.")
        cols=["id","tx_date","description","amount_mills","date_status","source_uid","source_period",
              "source_cell","source_order","created_at","updated_at"]
        with self.conn:
            self.conn.execute(f"INSERT INTO transactions({','.join(cols)}) VALUES({','.join('?'*len(cols))})",
                              [tx[c] for c in cols])
            if snapshot.get("comment"):
                self.conn.execute("INSERT INTO transaction_comments(transaction_id,note) VALUES(?,?)",
                                  (tx["id"],snapshot["comment"]))
            if snapshot.get("category"):
                self.conn.execute("INSERT INTO transaction_categories(transaction_id,category) VALUES(?,?)",
                                  (tx["id"],snapshot["category"]))
            for post in snapshot.get("recurring_posts",[]):
                # Restore the post record only if its template still exists and the month
                # was not posted again in the meantime.
                self.conn.execute("""INSERT OR IGNORE INTO recurring_posts(template_id,month,transaction_id)
                    SELECT ?,?,? WHERE EXISTS (SELECT 1 FROM recurring_templates WHERE id=?)""",
                    (post["template_id"],post["month"],tx["id"],post["template_id"]))
            for loan_id in snapshot.get("dismissed",[]):
                self.conn.execute("""INSERT OR IGNORE INTO dismissed_suggestions(repayment_tx_id,loan_id)
                    SELECT ?,? WHERE EXISTS (SELECT 1 FROM loan_records WHERE id=?)""",(tx["id"],loan_id,loan_id))
        return int(tx["id"])

    def entries(self) -> list[Entry]:
        # Imported source order is used ONLY for same-day tie-breaking; chronological
        # date is always primary. Manually entered transactions sort after imported
        # historical records for the same date.
        raw=self.conn.execute("""SELECT * FROM transactions ORDER BY tx_date ASC,
            CASE WHEN source_order IS NULL THEN 1 ELSE 0 END ASC,
            source_order ASC, id ASC""").fetchall()
        balance=self.opening()[0]
        result=[]
        for row in raw:
            balance+=int(row["amount_mills"])
            result.append(Entry(
                int(row["id"]),row["tx_date"],row["description"],int(row["amount_mills"]),
                row["date_status"],row["source_uid"],row["source_period"],
                row["source_cell"],row["source_order"],balance
            ))
        return result

    def balance(self) -> int:
        row=self.conn.execute("SELECT COALESCE(SUM(amount_mills),0) FROM transactions").fetchone()
        return self.opening()[0]+int(row[0])

    def monthly_totals(self) -> list[dict]:
        entries=self.entries()
        result=[]
        months={}
        for e in entries:
            month=e.date[:7]
            if month not in months:
                months[month]={"month": month, "income":0, "expense":0, "count":0,
                               "closing":e.running_mills}
                result.append(months[month])
            row=months[month]
            row["count"]+=1
            row["income"]+=max(e.mills,0)
            row["expense"]+=max(-e.mills,0)
            row["closing"]=e.running_mills
        return result

    def import_data(self, data) -> tuple[int,int]:
        """Atomic import with source-UID deduplication and opening-balance protection."""
        seen=set()
        for tx in data.transactions:
            if tx.source_uid in seen:
                raise ValueError("Import file contains duplicate source IDs.")
            seen.add(tx.source_uid)
            self._check_entry(tx.date,tx.description,tx.mills,tx.date_status)
            categories.clean(getattr(tx,"category",""))
        current_opening=self.opening()[0]
        if data.opening_mills is not None and (self.count()>0 or current_opening != 0) and current_opening != data.opening_mills:
            raise ValueError("Opening balance differs from your existing ledger. Import canceled to protect your data.")
        existing={r[0] for r in self.conn.execute("SELECT source_uid FROM transactions")}
        inserted=skipped=0
        with self.conn:
            if data.opening_mills is not None and self.count()==0:
                self.conn.execute("UPDATE state SET opening_mills=?,opening_date=?,opening_status=? WHERE id=1",
                                  (data.opening_mills,data.opening_date or '1970-01-01',data.opening_status))
            now=utc_now()
            for tx in data.transactions:
                if tx.source_uid in existing:
                    skipped+=1
                    continue
                cur=self.conn.execute("""INSERT INTO transactions (tx_date,description,amount_mills,
                    date_status,source_uid,source_period,source_cell,source_order,created_at,updated_at)
                    VALUES (?,?,?,?,?,?,?,?,?,?)""", (
                    tx.date, tx.description, tx.mills, tx.date_status, tx.source_uid,
                    tx.source_period, tx.source_cell, tx.source_order, now,now
                ))
                category=categories.clean(getattr(tx,"category",""))
                if category:
                    self.conn.execute("INSERT INTO transaction_categories(transaction_id,category) VALUES(?,?)",
                                      (cur.lastrowid,category))
                inserted+=1
        return inserted,skipped

    def duplicates_in(self, data) -> int:
        existing={r[0] for r in self.conn.execute("SELECT source_uid FROM transactions")}
        return sum(1 for tx in data.transactions if tx.source_uid in existing)

    def export_csv(self, path: str | Path) -> None:
        fields=["Date","Description","Amount (+/-)","Running balance","Entry type",
                "Date status","Date rule","Source period","Source cell","Source order","Stable import ID","Category"]
        opening_mills, opening_date, opening_status=self.opening()
        cats=categories.mapping(self.conn)
        with open(path,"w",newline="",encoding="utf-8-sig") as fp:
            writer=csv.writer(fp)
            writer.writerow(fields)
            writer.writerow([opening_date,"Opening balance (not income)",mills_as_decimal(opening_mills),
                             mills_as_decimal(opening_mills),"Opening balance",opening_status.title(),
                             "User-provided balance","","",0,"glassbalance:opening",""])
            for e in self.entries():
                writer.writerow([e.date,e.description,mills_as_decimal(e.mills),
                                 mills_as_decimal(e.running_mills),"Transaction",e.date_status.title(),
                                 "","" if not e.source_period else e.source_period,
                                 e.source_cell,e.source_order if e.source_order is not None else "",e.source_uid,
                                 cats.get(e.id,"")])

    def export_xlsx(self, path: str | Path) -> None:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.table import Table, TableStyleInfo
        opening_mills, opening_date, opening_status=self.opening()
        headers=["Date","Description","Amount (+/-)","Running balance","Entry type",
                 "Date status","Date rule","Source period","Source cell","Source order","Stable import ID","Category"]
        cats=categories.mapping(self.conn)
        wb=Workbook()
        ws=wb.active
        ws.title="Transactions"
        ws.append(headers)
        ws.append([date.fromisoformat(opening_date),"Opening balance (not income)",opening_mills/1000,
                   opening_mills/1000,"Opening balance",opening_status.title(),"User-provided balance","","",0,"glassbalance:opening",""])
        entries=self.entries()
        for row,e in enumerate(entries,start=3):
            ws.append([date.fromisoformat(e.date),e.description,e.mills/1000,
                       f"=D{row-1}+C{row}","Transaction",e.date_status.title(),"",
                       e.source_period,e.source_cell,e.source_order,e.source_uid,cats.get(e.id,"")])
        for cell in ws[1]:
            cell.fill=PatternFill("solid",fgColor="26364C")
            cell.font=Font(color="FFFFFF",bold=True)
            cell.alignment=Alignment(wrap_text=True)
        ws.row_dimensions[1].height=30
        for i in range(2,ws.max_row+1):
            ws[f"A{i}"].number_format="yyyy-mm-dd"
            if i>2 and entries[i-3].mills % 10:
                ws[f"C{i}"].number_format="$#,##0.000;[Red]($#,##0.000);-"
            for col in "CD":
                if not (col=="C" and i>2 and entries[i-3].mills % 10):
                    ws[f"{col}{i}"].number_format='$#,##0.00;[Red]($#,##0.00);-'
        for col,w in {"A":16,"B":35,"C":19,"D":19,"E":20,"F":18,"G":20,"H":32,"I":17,"J":17,"K":45,"L":22}.items():
            ws.column_dimensions[col].width=w
        ws.freeze_panes="C3"
        ws.sheet_view.showGridLines=False
        ws.auto_filter.ref=f"A1:L{ws.max_row}"
        wb.calculation.fullCalcOnLoad=True
        wb.save(path)

    def backup_to(self, path: str | Path) -> None:
        path=Path(path)
        if self.path and path.resolve() == self.path.resolve():
            raise ValueError("Choose a backup file different from the active database.")
        path.parent.mkdir(parents=True,exist_ok=True)
        target=sqlite3.connect(str(path))
        try:
            self.conn.backup(target)
        finally:
            target.close()

    def backup_bytes(self) -> bytes:
        with tempfile.TemporaryDirectory() as d:
            file=Path(d)/"backup.sqlite"
            self.backup_to(file)
            return file.read_bytes()

    def restore_bytes(self, payload: bytes) -> None:
        with tempfile.TemporaryDirectory() as d:
            file=Path(d)/"restore.sqlite"
            file.write_bytes(payload)
            self.restore_from(file)

    def restore_from(self, path: str | Path) -> None:
        """Validate an SQLite backup, then atomically replace this DB's content.

        UI must request explicit confirmation and take an independent local backup first.
        """
        if self.path and Path(path).resolve() == self.path.resolve():
            raise ValueError("Cannot restore the active database onto itself.")
        source=sqlite3.connect(Path(path).resolve().as_uri()+"?mode=ro", uri=True)
        try:
            integrity=source.execute("PRAGMA integrity_check").fetchone()
            if not integrity or integrity[0]!="ok":
                raise ValueError("Backup failed its SQLite integrity check.")
            required={"state","transactions"}
            tables={r[0] for r in source.execute("SELECT name FROM sqlite_master WHERE type='table'")}
            if not required.issubset(tables):
                raise ValueError("This is not a compatible GlassBalance backup.")
            source.execute("SELECT opening_mills FROM state WHERE id=1").fetchone()
            source.execute("SELECT amount_mills,source_uid FROM transactions LIMIT 1").fetchall()
            source.backup(self.conn)
        finally:
            source.close()
        # Older backups may predate comments, loans, categories or recurring tables.
        self._install_extensions()
