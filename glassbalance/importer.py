"""Import curated CSV/XLSX or the user's original horizontal balance workbook.

The original never had transaction dates. Generated dates are clearly flagged
'estimated' and preserve the original source cell for audits/deduplication.
"""
from __future__ import annotations

import calendar
import csv
import hashlib
import re
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path

from .money import as_mills


@dataclass(frozen=True)
class ImportTransaction:
    date: str
    description: str
    mills: int
    date_status: str
    source_uid: str
    source_period: str = ""
    source_cell: str = ""
    source_order: int | None = None
    category: str = ""


@dataclass
class ImportData:
    transactions: list[ImportTransaction]
    opening_mills: int | None = None
    opening_date: str | None = None
    opening_status: str = "estimated"
    format: str = ""
    note: str = ""


MONTHS={"jan":1,"feb":2,"mar":3,"apr":4,"may":5,"jun":6,"jul":7,
        "aug":8,"sep":9,"oct":10,"nov":11,"dec":12}
MONTH_RE=re.compile(r"\b(january|jan|february|feb|march|mar|april|apr|may|june|jun|"
                    r"july|jul|august|aug|september|sept|sep|october|oct|november|nov|december|dec)\b",re.I)
SALARY_RE=re.compile(r"\bsalary\b",re.I)
RENT_RE=re.compile(r"\brent\b",re.I)


def _month_from_text(text: str) -> int | None:
    matched=MONTH_RE.search(text)
    return MONTHS[matched.group(1).lower()[:3]] if matched else None


def _date_value(value: object) -> str:
    if isinstance(value,datetime):
        return value.date().isoformat()
    if isinstance(value,date):
        return value.isoformat()
    value=str(value).strip()
    try:
        return date.fromisoformat(value).isoformat()
    except ValueError as exc:
        raise ValueError(f"Expected dates as YYYY-MM-DD, got {value!r}.") from exc


def _rent(title: str, amount: int) -> bool:
    if amount >= 0 or not RENT_RE.search(title):
        return False
    lowered=title.lower()
    if "refund" in lowered or re.search(r"\bcar\s+rent\b|\brent\s+tax\b",lowered):
        return False
    return True


def _next_month(year: int, month: int) -> tuple[int,int]:
    return (year+1,1) if month==12 else (year,month+1)


def _parsed_original(ws, as_of: date) -> ImportData:
    # Original source: repeated 'Name | Bill or Income | Amount | Unbalanced'
    # sections running left-to-right in one worksheet.
    heading_columns=[]
    for col in range(1,ws.max_column):
        name=ws.cell(1,col).value
        amt=ws.cell(1,col+1).value
        if isinstance(name,str) and name.strip().lower()=="name" and isinstance(amt,str) and amt.strip().lower()=="bill or income":
            heading_columns.append(col)
    if not heading_columns:
        raise ValueError("Can't find the source workbook's Name / Bill or Income sections.")
    opening=ws["C2"].value
    if not isinstance(opening,(int,float)) or isinstance(opening,bool):
        raise ValueError("Original workbook opening balance must be numeric in C2.")
    groups=[]
    first_year=None
    for col in heading_columns:
        period=""
        for possible in range(col+3,min(col+8,ws.max_column+1)):
            val=ws.cell(2,possible).value
            if isinstance(val,str) and re.search(r"\b20\d{2}\b",val):
                period=val.strip()
                break
        if first_year is None and period:
            year=re.search(r"\b(20\d{2})\b",period)
            if year:
                first_year=int(year.group(1))
        groups.append((col,period))
    first_year=first_year or 2022  # explicit historical default, flagged estimated
    records=[]
    for col,period in groups:
        for row in range(3,ws.max_row+1):
            value=ws.cell(row,col+1).value
            if isinstance(value,bool) or not isinstance(value,(float,int)):
                continue
            name=ws.cell(row,col).value
            if name is None or not str(name).strip():
                raise ValueError(f"Missing description in {ws.cell(row,col).coordinate}")
            records.append({
                "description":str(name).strip(),"mills":as_mills(value),
                "source_period":period,"source_cell":ws.cell(row,col).coordinate,
                "source_order":len(records)+1,
                "source_uid":f"personal-balance-ledger:{ws.title}!{ws.cell(row,col).coordinate}"
            })
    if not records:
        raise ValueError("Found sections, but no transactions.")
    payroll_year=first_year
    last_payroll_month=0
    cycles={}
    for r in records:
        salary=SALARY_RE.search(r["description"]) is not None and r["mills"]>0
        if salary:
            month=_month_from_text(r["description"])
            if month is None:
                raise ValueError(f"Cannot identify salary month in {r['source_cell']}: {r['description']}")
            if month<last_payroll_month:
                payroll_year+=1
            last_payroll_month=month
            next_year,next_month=_next_month(payroll_year,month)
            current_cycle=date(next_year,next_month,1)
            r["estimated_date"]=date(payroll_year,month,calendar.monthrange(payroll_year,month)[1])
            r["rule"]="salary"
        else:
            if last_payroll_month==0:
                raise ValueError("Historical dates require a salary to anchor the first cycle.")
            if _rent(r["description"],r["mills"]):
                r["estimated_date"]=current_cycle.replace(day=2)
                r["rule"]="rent"
            else:
                r["rule"]="other"
        r["cycle"]=current_cycle
        cycles.setdefault(current_cycle,[]).append(r)

    for cycle,items in cycles.items():
        days=calendar.monthrange(cycle.year,cycle.month)[1]
        if (cycle.year,cycle.month)==(as_of.year,as_of.month):
            days=min(days,as_of.day)
        if (cycle.year,cycle.month)>(as_of.year,as_of.month):
            raise ValueError("Original contains an estimated pay-cycle later than today's date.")
        days=max(days,3)
        ordinary=[r for r in items if r["rule"]=="other"]
        n=len(ordinary)
        for i,r in enumerate(ordinary):
            day=3+(int((days-3)*i/(n-1)+0.5) if n>1 else (days-3)//2)
            r["estimated_date"]=cycle.replace(day=day)

    # Handle a bookkeeping boundary with two rent debits immediately before
    # the next month's salary while the following month contains no rent.
    for cycle,items in cycles.items():
        following=date(*_next_month(cycle.year,cycle.month),1)
        if following not in cycles:
            continue
        rents=[r for r in items if r["rule"]=="rent"]
        next_rents=[r for r in cycles[following] if r["rule"]=="rent"]
        following_salary=[r["source_order"] for r in cycles[following] if r["rule"]=="salary"]
        if len(rents)<2 or next_rents or not following_salary:
            continue
        before=[r for r in rents if 0<min(following_salary)-r["source_order"]<=2]
        if before:
            max(before,key=lambda x:x["source_order"])["estimated_date"]=following.replace(day=2)

    ordered=sorted(records,key=lambda r:(r["estimated_date"],r["source_order"]))
    first=min(r["estimated_date"] for r in ordered)
    output=[ImportTransaction(r["estimated_date"].isoformat(),r["description"],r["mills"],
                              "estimated",r["source_uid"],r["source_period"],r["source_cell"],
                              r["source_order"]) for r in ordered]
    return ImportData(output,as_mills(opening),(first-timedelta(days=1)).isoformat(),
                      "estimated","Original horizontal workbook",
                      "Dates are estimates: salary is assigned to the last day of the named month, rent to the next month's 2nd, and remaining entries spread across the next month.")


def _column_dictionary(row) -> dict:
    """Use predictable case-insensitive headers while allowing whitespace."""
    return {str(key).strip().lower():value for key,value in row.items() if key is not None}


def _from_rows(rows: list[dict], file_fingerprint: str, fmt: str) -> ImportData:
    if not rows:
        raise ValueError("The sheet is empty.")
    example=_column_dictionary(rows[0])
    if not any(c in example for c in ("amount (+/-)","amount","bill or income")):
        raise ValueError("Expected an Amount (+/-) or Amount column.")
    if not any(c in example for c in ("date","transaction date")):
        raise ValueError("Expected a Date column.")
    if not any(c in example for c in ("description","name")):
        raise ValueError("Expected a Description or Name column.")
    imported=[]
    opening=None
    opening_date=None
    opening_status="estimated"
    for index,row in enumerate(rows,start=2):
        r=_column_dictionary(row)
        raw_desc=r.get("description",r.get("name"))
        if raw_desc is None or str(raw_desc).strip()=="":
            if not any(v is not None and str(v).strip() for v in r.values()):
                continue
            raise ValueError(f"Missing description in row {index}.")
        desc=str(raw_desc).strip()
        amount=r.get("amount (+/-)",r.get("amount",r.get("bill or income")))
        if amount is None or str(amount).strip()=="":
            raise ValueError(f"Missing amount in row {index}.")
        mills=as_mills(amount)
        when=_date_value(r.get("date",r.get("transaction date")))
        status=str(r.get("date status") or "Confirmed").strip().lower()
        if status not in ("estimated","confirmed"):
            raise ValueError(f"Date status must be Estimated or Confirmed (row {index}).")
        kind=str(r.get("entry type") or "Transaction").strip().lower()
        if kind=="opening balance":
            if opening is not None:
                raise ValueError("Input includes multiple opening-balance rows.")
            opening=mills
            opening_date=when
            opening_status=status
            continue
        if kind not in ("transaction", ""):
            raise ValueError(f"Unrecognized Entry type in row {index}: {kind}")
        source_id=str(r.get("stable import id") or "").strip()
        if not source_id:
            source_id=f"glassbalance:import:{file_fingerprint}:{index}"
        raw_order=r.get("source order")
        order=None if raw_order is None or str(raw_order).strip()=="" else int(raw_order)
        imported.append(ImportTransaction(when,desc,mills,status,source_id,
                        str(r.get("source period") or ""),str(r.get("source cell") or ""),order,
                        str(r.get("category") or "").strip()))
    return ImportData(imported,opening,opening_date,opening_status,fmt,
                      "Estimated dates are carried through from the imported sheet. Manual edits in the app will mark dates as confirmed.")


def load_import(path: str | Path, as_of: date | None = None) -> ImportData:
    """Read an input file WITHOUT changing or saving the source file."""
    path=Path(path)
    if not path.is_file():
        raise FileNotFoundError(path)
    if path.suffix.lower() not in (".xlsx",".csv"):
        raise ValueError("Import supports .xlsx and .csv only.")
    sha=hashlib.sha256(path.read_bytes()).hexdigest()[:24]
    if path.suffix.lower()==".csv":
        with path.open("r",newline="",encoding="utf-8-sig") as fp:
            rows=list(csv.DictReader(fp))
        return _from_rows(rows,sha,"CSV")
    from openpyxl import load_workbook
    wb=load_workbook(path,read_only=False,data_only=True)
    try:
        ws=wb["Transactions"] if "Transactions" in wb.sheetnames else wb.worksheets[0]
        # Curation produces a standard Transactions sheet; original workbook
        # contains separate horizontal blocks and uses a different parser.
        if ws.max_row and str(ws.cell(1,1).value or "").strip().lower()=="date":
            header=[str(c.value).strip() if c.value is not None else "" for c in ws[1]]
            rows=[dict(zip(header,row)) for row in ws.iter_rows(min_row=2,values_only=True)]
            return _from_rows(rows,sha,"Curated/standard Excel")
        return _parsed_original(ws,as_of or date.today())
    finally:
        wb.close()
